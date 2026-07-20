"""DOCX exporter via `python-docx`. No round-trip editing fidelity is claimed -
this is one-way: compile IR in, a plain Word document out."""

from __future__ import annotations

from pathlib import Path

from docx import Document as create_docx_document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH

from skrybe_studio.plugins.manuscript.application.export.compile import (
    CompileDocument,
    CompileSection,
)
from skrybe_studio.plugins.manuscript.domain.document_ir import DocumentBlock
from skrybe_studio.plugins.manuscript.domain.enums import ManuscriptNodeType

SCENE_SEPARATOR = "* * *"
_STYLE_BY_BLOCK_KIND = {"blockquote": "Quote", "list_item": "List Bullet"}


def export_docx(document: CompileDocument, output_path: Path) -> Path:
    docx = create_docx_document()
    docx.add_heading(document.title, level=0)
    _walk(document.sections, docx)
    docx.save(str(output_path))
    return output_path


def _walk(sections: tuple[CompileSection, ...], docx: DocxDocument) -> None:
    previous_was_text = False
    for section in sections:
        is_text = section.node_type in (ManuscriptNodeType.SCENE, ManuscriptNodeType.DOCUMENT)
        if is_text:
            if previous_was_text:
                separator = docx.add_paragraph(SCENE_SEPARATOR)
                separator.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for block in section.blocks:
                _add_block(docx, block)
        else:
            docx.add_heading(section.title, level=min(max(section.level, 1), 9))
            _walk(section.children, docx)
        previous_was_text = is_text


def _add_block(docx: DocxDocument, block: DocumentBlock) -> None:
    style = _STYLE_BY_BLOCK_KIND.get(block.kind.value)
    paragraph = docx.add_paragraph(style=style) if style else docx.add_paragraph()
    for span in block.spans:
        run = paragraph.add_run(span.text)
        run.bold = span.bold
        run.italic = span.italic
        run.underline = span.underline
